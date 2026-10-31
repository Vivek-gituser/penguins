"""Real PowerPoint output built with python-pptx from a validated DocumentSpec.

The deck always has a title slide plus one slide per section, a consistent
theme, and a classification footer drawn from the inherited source label.
"""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

from tools.document_spec import DocumentSpec

PRIMARY = RGBColor(0x4F, 0x46, 0xE5)
ACCENT = RGBColor(0x14, 0xB8, 0xA6)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xEE, 0xF2, 0xFF)

LABEL_COLORS = {
    "PUBLIC": RGBColor(0x04, 0x78, 0x57),
    "INTERNAL": RGBColor(0x1D, 0x4E, 0xD8),
    "CONFIDENTIAL": RGBColor(0xB4, 0x53, 0x09),
    "RESTRICTED": RGBColor(0xB9, 0x1C, 0x1C),
}

SLIDE_WIDTH = Emu(12192000)   # 13.333 in
SLIDE_HEIGHT = Emu(6858000)   # 7.5 in
MARGIN = Emu(685800)          # 0.75 in
FOOTER_HEIGHT = Emu(274320)   # 0.3 in


def _add_slide(presentation):
    """Blank layout, so the theme never depends on a template placeholder."""
    for layout in presentation.slide_layouts:
        if layout.name.lower().startswith("blank"):
            return presentation.slides.add_slide(layout)
    return presentation.slides.add_slide(presentation.slide_layouts[6])


def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _frame(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = box.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = anchor
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    return text_frame


def _write(text_frame, lines, size, color, bold=False, space_after=8):
    for index, text in enumerate(lines):
        paragraph = (
            text_frame.paragraphs[0] if index == 0
            else text_frame.add_paragraph()
        )
        paragraph.level = 0
        paragraph.space_after = Pt(space_after)
        paragraph.alignment = PP_ALIGN.LEFT

        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Calibri"
        run.font.color.rgb = color


def _bullets(text_frame, bullets, size=18, space_after=12):
    for index, bullet in enumerate(bullets):
        paragraph = (
            text_frame.paragraphs[0] if index == 0
            else text_frame.add_paragraph()
        )
        paragraph.space_after = Pt(space_after)
        paragraph.alignment = PP_ALIGN.LEFT

        marker = paragraph.add_run()
        marker.text = "\u2022  "
        marker.font.size = Pt(size)
        marker.font.bold = True
        marker.font.name = "Calibri"
        marker.font.color.rgb = PRIMARY

        body = paragraph.add_run()
        body.text = bullet
        body.font.size = Pt(size)
        body.font.name = "Calibri"
        body.font.color.rgb = INK


def _footer(slide, spec, page_number, total):
    color = LABEL_COLORS.get(spec.label, LABEL_COLORS["INTERNAL"])
    top = SLIDE_HEIGHT - FOOTER_HEIGHT

    _rect(slide, 0, top, Emu(2743200), FOOTER_HEIGHT, color)

    _write(
        _frame(slide, Emu(228600), top, Emu(2400000), FOOTER_HEIGHT,
               MSO_ANCHOR.MIDDLE),
        [spec.label], 11, PAPER, bold=True,
    )
    _write(
        _frame(slide, SLIDE_WIDTH - MARGIN - Emu(3600000), top,
               Emu(3600000), FOOTER_HEIGHT, MSO_ANCHOR.MIDDLE),
        [f"AGENTGATE workspace  |  slide {page_number} of {total}"],
        10, MUTED,
    )


def build_pptx(spec):
    """Return genuine .pptx bytes for the given spec. Never writes a stub."""
    if not isinstance(spec, DocumentSpec):
        raise ValueError("A validated DocumentSpec is required")

    if not spec.sections:
        raise ValueError("A presentation needs at least one content section")

    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    total = spec.slide_count
    content_width = SLIDE_WIDTH - MARGIN - Emu(914400)

    # ---- title slide ----
    cover = _add_slide(presentation)
    _rect(cover, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, PAPER)
    _rect(cover, 0, 0, MARGIN, SLIDE_HEIGHT, PRIMARY)
    _rect(cover, MARGIN, SLIDE_HEIGHT - Emu(1600200),
          SLIDE_WIDTH - MARGIN, Emu(1600200), SOFT)

    _write(
        _frame(cover, MARGIN + Emu(457200), Emu(1500000), content_width,
               Emu(2400000)),
        [spec.title], 40, PRIMARY, bold=True, space_after=10,
    )

    if spec.subtitle:
        _write(
            _frame(cover, MARGIN + Emu(457200), Emu(3900000), content_width,
                   Emu(900000)),
            [spec.subtitle], 18, MUTED,
        )

    _write(
        _frame(cover, MARGIN + Emu(457200), SLIDE_HEIGHT - Emu(1350000),
               content_width, Emu(1000000)),
        [
            f"Security classification: {spec.label}",
            f"{len(spec.sections)} content slides - prepared with AGENTGATE",
        ],
        14, INK, space_after=4,
    )
    _footer(cover, spec, 1, total)

    # ---- section slides ----
    for index, section in enumerate(spec.sections, start=1):
        slide = _add_slide(presentation)
        _rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, PAPER)
        _rect(slide, 0, 0, SLIDE_WIDTH, Emu(1100000), PRIMARY)
        _rect(slide, 0, Emu(1100000), SLIDE_WIDTH, Emu(70000), ACCENT)

        _write(
            _frame(slide, MARGIN, Emu(160000),
                   SLIDE_WIDTH - MARGIN - Emu(914400), Emu(800000),
                   MSO_ANCHOR.MIDDLE),
            [section.heading], 24, PAPER, bold=True,
        )
        _write(
            _frame(slide, SLIDE_WIDTH - MARGIN - Emu(600000), Emu(160000),
                   Emu(600000), Emu(800000), MSO_ANCHOR.MIDDLE),
            [f"{index:02d}"], 22, SOFT, bold=True,
        )

        _bullets(
            _frame(slide, MARGIN, Emu(1500000), SLIDE_WIDTH - MARGIN * 2,
                   SLIDE_HEIGHT - Emu(2000000)),
            section.bullets,
        )

        _footer(slide, spec, index + 1, total)

    buffer = io.BytesIO()
    presentation.save(buffer)
    payload = buffer.getvalue()

    if not payload.startswith(b"PK") or len(payload) < 2000:
        raise ValueError("PowerPoint generation did not produce a real file")

    return payload
